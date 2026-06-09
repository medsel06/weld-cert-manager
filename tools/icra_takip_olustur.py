# -*- coding: utf-8 -*-
"""
Personel İcra Takip Excel şablonu üretici (v2 - makro destekli).

Zirve Müşavir "Personel İcra Bilgi Girişi" ekranındaki mantığı Excel'e taşır:
  - PERSONEL sayfasında net maaş tutulur.
  - Her icra dosyası öncelik sırasıyla (1, 2, 3...) izlenir; kesinti kuralı
    Zirve'deki gibi bir kez girilir: Oran (pay/payda, örn. 1/4) veya Sabit Tutar.
  - "Bu Ay Kesilecek" kolonu kesintiyi kendiliğinden hesaplar
    (net maaş x oran, kalan borçla sınırlı, yalnız ÖDEMEDE'deki dosya).
  - PANEL sayfasından yıl/ay seçilip VBA makrosuyla (vba/IcraTakip.bas)
    "KAYDET" denince kesintiler AYLIK KESİNTİLER'e otomatik işlenir.
  - Kalan borç 0'a inen dosya KAPANDI olur, sıradaki öncelik otomatik
    ÖDEMEDE'ye geçer; aynı ay içinde artan tutar sonraki dosyaya taşar.

Kullanım:  python3 tools/icra_takip_olustur.py [çıktı.xlsx]
"""
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

PERSONEL_SON = 300
DOSYA_SON_SATIR = 300
KESINTI_SON_SATIR = 5000
OZET_SON_SATIR = 100

LACIVERT = "1F4E78"   # elle girilen kolon başlığı
TURUNCU = "C65911"    # otomatik (formüllü) kolon başlığı
YESIL = "548235"

INCE = Side(style="thin", color="BFBFBF")
KENARLIK = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
PARA = "#,##0.00"
TARIH = "DD.MM.YYYY"
AYLAR = "Ocak,Şubat,Mart,Nisan,Mayıs,Haziran,Temmuz,Ağustos,Eylül,Ekim,Kasım,Aralık"


def baslik_yaz(ws, basliklar, otomatikler=()):
    for i, ad in enumerate(basliklar, start=1):
        c = ws.cell(row=1, column=i, value=ad)
        renk = TURUNCU if i in otomatikler else LACIVERT
        c.fill = PatternFill("solid", start_color=renk, end_color=renk)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = KENARLIK
    ws.row_dimensions[1].height = 34


def genislik(ws, genislikler):
    for i, w in enumerate(genislikler, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def durum_boyamasi(ws, aralik):
    for metin, zemin, yazi in [("ÖDEMEDE", "C6EFCE", "006100"),
                               ("SIRADA", "FFEB9C", "9C6500"),
                               ("KAPANDI", "D9D9D9", "595959")]:
        ws.conditional_formatting.add(
            aralik,
            CellIsRule(operator="equal", formula=[f'"{metin}"'],
                       fill=PatternFill("solid", start_color=zemin, end_color=zemin),
                       font=Font(bold=True, color=yazi)))


def kullanim_sayfasi(wb):
    ws = wb.active
    ws.title = "KULLANIM"
    ws.sheet_properties.tabColor = "808080"
    ws.column_dimensions["A"].width = 120

    satirlar = [
        ("PERSONEL İCRA TAKİP DOSYASI (makro destekli)", "baslik"),
        ("Bir personelin birden fazla icra dosyasını öncelik sırasına göre izler; net maaştan oranlı (1/4) "
         "veya sabit tutarlı kesintiyi kendisi hesaplar, 'KAYDET' butonuyla aylık kesintileri otomatik işler.", None),
        ("", None),
        ("SAYFALAR", "baslik2"),
        ("1) PERSONEL: Her personeli bir kez girin: Ad Soyad, TC, NET MAAŞ. Maaş değişince burada güncelleyin.", None),
        ("2) İCRA DOSYALARI: Her icra dosyası AYRI SATIR (5 icrası olan personel için 5 satır). "
         "'Öncelik Sırası'na 1, 2, 3... yazın. Kesinti kuralını BİR KEZ girin: "
         "Kesinti Şekli = Oran ise Pay/Payda kutularına örn. 1 ve 4 yazın (Zirve'deki 1/4 gibi); "
         "Sabit Tutar ise aylık tutarı yazın. TC, net maaş, tahsilat, kalan borç ve DURUM otomatiktir.", None),
        ("3) PANEL: Ay sonunda yıl/ay seçin, 'KESİNTİLERİ HESAPLA ve KAYDET' butonuna basın. "
         "Makro, ÖDEMEDE durumundaki dosyalara o ayın kesintisini hesaplayıp AYLIK KESİNTİLER'e işler "
         "(açıklamasına OTOMATİK yazar). Aynı ay ikinci kez basarsanız mükerrer kayıt OLUŞMAZ, atlanır.", None),
        ("4) AYLIK KESİNTİLER: Kesinti dökümü buraya birikir. Elle de satır girebilirsiniz "
         "(örn. geçmiş aylar). 'Geri Al' butonu yalnız seçili ayın OTOMATİK satırlarını siler.", None),
        ("5) PERSONEL ÖZET: Personel bazında toplamlar, ödemedeki dosya ve bu ay kesilecek tutar.", None),
        ("", None),
        ("HESAP MANTIĞI", "baslik2"),
        ("Bu Ay Kesilecek = Net Maaş x Pay/Payda (veya Sabit Tutar); kalan borcu aşamaz. "
         "Kalan borcu biten dosya KAPANDI olur, sıradaki öncelik OTOMATİK ÖDEMEDE'ye geçer. "
         "Kaydederken dosya kapanır da tutar artarsa, artan kısım aynı ay sıradaki dosyaya işlenir.", None),
        ("", None),
        ("MAKRO KURULUMU (bir kez, ~1 dakika)", "baslik2"),
        ("ADIM 0: Dosyalar internetten/WhatsApp'tan geldiyse: bu Excel'e sağ tık > Özellikler > "
         "en altta 'Engellemeyi Kaldır' (Unblock) işaretli ise işaretleyip Tamam deyin.", None),
        ("ADIM 1: Bu dosyayı Excel'de açın, Alt+F11 tuşuna basın (VBA editörü açılır).", None),
        ("ADIM 2: Üst menüden Dosya > Dosya Al... (File > Import File) deyip IcraTakip.bas dosyasını seçin.", None),
        ("ADIM 3: Alt+Q ile editörden çıkın. Dosya > Farklı Kaydet > kayıt türü olarak "
         "'Excel Makro İçerebilen Çalışma Kitabı (*.xlsm)' seçip kaydedin.", None),
        ("ADIM 4: Dosyayı kapatıp .xlsm olanı açın; sarı çubukta 'İçeriği Etkinleştir' çıkarsa tıklayın.", None),
        ("ADIM 5: Alt+F8 > BUTONLARI_KUR > Çalıştır. Butonlar PANEL sayfasına eklenir. Kurulum bitti.", None),
        ("", None),
        ("RENK REHBERİ", "baslik2"),
        ("LACİVERT başlıklı kolonlar elle doldurulur.", "lacivert"),
        ("TURUNCU başlıklı kolonlar formüllüdür; elle veri GİRMEYİN, silmeyin.", "turuncu"),
        ("DURUM RENKLERİ:   YEŞİL = ÖDEMEDE     SARI = SIRADA     GRİ = KAPANDI", None),
        ("Formüller 300 dosya / 5.000 kesinti satırı için hazırdır; gerekirse son formüllü satırı aşağı kopyalayın.", None),
        ("Örnek kayıtların açıklamasında 'ÖRNEK' yazar; kendi verinizi girerken silebilirsiniz.", None),
    ]
    for r, (metin, stil) in enumerate(satirlar, start=2):
        c = ws.cell(row=r, column=1, value=metin)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if stil == "baslik":
            c.font = Font(bold=True, size=14, color=LACIVERT)
        elif stil == "baslik2":
            c.font = Font(bold=True, size=11, color=LACIVERT)
        elif stil == "lacivert":
            c.fill = PatternFill("solid", start_color=LACIVERT, end_color=LACIVERT)
            c.font = Font(bold=True, color="FFFFFF")
        elif stil == "turuncu":
            c.fill = PatternFill("solid", start_color=TURUNCU, end_color=TURUNCU)
            c.font = Font(bold=True, color="FFFFFF")
        ws.row_dimensions[r].height = 42 if len(metin) > 110 else (30 if len(metin) > 60 else 18)


def personel_sayfasi(wb):
    ws = wb.create_sheet("PERSONEL")
    ws.sheet_properties.tabColor = "7030A0"
    baslik_yaz(ws, ["Personel Adı Soyadı", "T.C. Kimlik No", "Net Maaş (TL)",
                    "İşe Giriş Tarihi", "Açıklama"])
    genislik(ws, [26, 15, 14, 13, 32])
    ws.freeze_panes = "A2"
    for r in range(2, PERSONEL_SON + 1):
        ws.cell(row=r, column=2).number_format = "@"
        ws.cell(row=r, column=3).number_format = PARA
        ws.cell(row=r, column=4).number_format = TARIH
        for kol in range(1, 6):
            ws.cell(row=r, column=kol).border = KENARLIK

    ornekler = [
        ["MERT TÜZE", "38128425238", 28075.52, datetime(2025, 3, 28), "ÖRNEK KAYIT"],
        ["AHMET BAŞOĞLU", "12856239602", 24000.00, datetime(2026, 2, 21), "ÖRNEK KAYIT"],
    ]
    for r, satir in enumerate(ornekler, start=2):
        for kol, deger in enumerate(satir, start=1):
            ws.cell(row=r, column=kol, value=deger)
    return ws


def icra_sayfasi(wb):
    icra = wb.create_sheet("İCRA DOSYALARI")
    icra.sheet_properties.tabColor = LACIVERT
    basliklar = [
        "Personel Adı Soyadı", "T.C. Kimlik No", "Net Maaş (TL)", "Dosya Tarihi",
        "Dosya Numarası", "Öncelik Sırası", "İcra Dairesi", "Alacaklı",
        "İcra Tutarı (TL)", "Kesinti Şekli", "Oran - Pay (örn. 1)", "Oran - Payda (örn. 4)",
        "Sabit Aylık Kesinti (TL)", "Bu Ay Kesilecek (TL)",
        "Önceki Firmada Tahsil (TL)", "Bordro Kesintileri Toplamı (TL)",
        "Toplam Tahsil Edilen (TL)", "Kalan Borç (TL)", "DURUM",
        "Ödeme Yapılacak Banka / IBAN", "Açıklama", "(otomatik - silmeyin)",
    ]
    baslik_yaz(icra, basliklar, otomatikler=(2, 3, 14, 16, 17, 18, 19, 22))
    genislik(icra, [24, 13, 12, 11, 15, 9, 28, 30, 13, 11, 9, 10, 12, 12, 12, 13, 13, 13, 12, 26, 26, 4])
    icra.freeze_panes = "B2"
    icra.auto_filter.ref = f"A1:U{DOSYA_SON_SATIR}"
    icra.column_dimensions["V"].hidden = True

    son = DOSYA_SON_SATIR
    for r in range(2, son + 1):
        icra.cell(row=r, column=2).value = (
            f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},PERSONEL!$A$2:$C${PERSONEL_SON},2,0),""))')
        icra.cell(row=r, column=3).value = (
            f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},PERSONEL!$A$2:$C${PERSONEL_SON},3,0),""))')
        icra.cell(row=r, column=14).value = (
            f'=IF($E{r}="","",IF($S{r}<>"ÖDEMEDE",0,'
            f'MIN($R{r},IF($J{r}="Oran",IFERROR($C{r}*$K{r}/$L{r},0),$M{r}))))')
        icra.cell(row=r, column=16).value = (
            f'=IF($E{r}="","",SUMIFS(\'AYLIK KESİNTİLER\'!$E$2:$E${KESINTI_SON_SATIR},'
            f"'AYLIK KESİNTİLER'!$A$2:$A${KESINTI_SON_SATIR},$A{r},"
            f"'AYLIK KESİNTİLER'!$B$2:$B${KESINTI_SON_SATIR},$E{r}))")
        icra.cell(row=r, column=17).value = f'=IF($E{r}="","",$O{r}+$P{r})'
        icra.cell(row=r, column=18).value = f'=IF($E{r}="","",$I{r}-$Q{r})'
        icra.cell(row=r, column=19).value = (
            f'=IF($E{r}="","",IF($R{r}<=0,"KAPANDI",'
            f'IF(COUNTIFS($A$2:$A${son},$A{r},$F$2:$F${son},"<"&$F{r},$R$2:$R${son},">0")=0,'
            f'"ÖDEMEDE","SIRADA")))')
        icra.cell(row=r, column=22).value = f'=IF($E{r}="","",$A{r}&"|"&$S{r})'
        for kol in (3, 9, 13, 14, 15, 16, 17, 18):
            icra.cell(row=r, column=kol).number_format = PARA
        icra.cell(row=r, column=2).number_format = "@"
        icra.cell(row=r, column=4).number_format = TARIH
        for kol in range(1, 22):
            icra.cell(row=r, column=kol).border = KENARLIK
        icra.cell(row=r, column=19).alignment = Alignment(horizontal="center")

    durum_boyamasi(icra, f"S2:S{son}")

    # Örnek kayıtlar: A,D,E,F,G,H,I,J,K,L,M,O,U
    ornekler = [
        ["MERT TÜZE", datetime(2026, 2, 23), "2025/203385", 1,
         "BANKA ALACAKLARI İCRA DAİRESİ", "AKBANK TÜRK ANONİM ŞİRKETİ",
         130656.65, "Oran", 1, 4, None, 0, "ÖRNEK (Zirve ekranındaki dosya)"],
        ["MERT TÜZE", datetime(2026, 6, 10), "deneme", 2,
         "", "", 10000.00, "Sabit Tutar", None, None, 10000.00, 0, "ÖRNEK"],
        ["AHMET BAŞOĞLU", datetime(2026, 1, 15), "2024/101010", 1,
         "ANKARA 5. İCRA DAİRESİ", "X FİNANS A.Ş.", 5000.00, "Sabit Tutar",
         None, None, 2500.00, 5000.00, "ÖRNEK: borç bitti, otomatik KAPANDI"],
        ["AHMET BAŞOĞLU", datetime(2026, 2, 21), "2025/55555", 2,
         "ANKARA 12. İCRA DAİRESİ", "Y BANKASI A.Ş.", 24000.00, "Oran", 1, 4,
         None, 0, "ÖRNEK: 1. dosya kapanınca ÖDEMEDE'ye geçti"],
        ["AHMET BAŞOĞLU", datetime(2026, 3, 1), "2026/77777", 3,
         "ANKARA 3. İCRA DAİRESİ", "Z TELEKOM A.Ş.", 8000.00, "Oran", 1, 4,
         None, 0, "ÖRNEK: sırada bekliyor"],
    ]
    hedef_kolonlar = (1, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 15, 21)
    for r, satir in enumerate(ornekler, start=2):
        for kol, deger in zip(hedef_kolonlar, satir):
            if deger is not None:
                icra.cell(row=r, column=kol, value=deger)

    dv_personel = DataValidation(type="list", formula1="PersonelAdlari", allow_blank=True)
    dv_personel.errorStyle = "warning"
    dv_personel.error = "PERSONEL sayfasında kayıtlı bir ad seçin (veya önce oraya ekleyin)."
    icra.add_data_validation(dv_personel)
    dv_personel.add(f"A2:A{son}")

    dv_sekil = DataValidation(type="list", formula1='"Oran,Sabit Tutar"', allow_blank=True)
    icra.add_data_validation(dv_sekil)
    dv_sekil.add(f"J2:J{son}")

    dv_oncelik = DataValidation(type="whole", operator="greaterThanOrEqual",
                                formula1="1", allow_blank=True)
    dv_oncelik.error = "1 veya daha büyük tam sayı girin (1 = ilk kesilecek dosya)."
    icra.add_data_validation(dv_oncelik)
    dv_oncelik.add(f"F2:F{son}")
    return icra


def kesinti_sayfasi(wb):
    kes = wb.create_sheet("AYLIK KESİNTİLER")
    kes.sheet_properties.tabColor = "2E75B6"
    baslik_yaz(kes, ["Personel Adı Soyadı", "Dosya Numarası", "Kesinti Yılı",
                     "Kesinti Dönemi (Ay)", "Kesinti Tutarı (TL)", "Açıklama"])
    genislik(kes, [24, 16, 11, 15, 15, 30])
    kes.freeze_panes = "A2"
    kes.auto_filter.ref = f"A1:F{KESINTI_SON_SATIR}"

    for r in range(2, KESINTI_SON_SATIR + 1):
        kes.cell(row=r, column=5).number_format = PARA
        if r <= 300:
            for kol in range(1, 7):
                kes.cell(row=r, column=kol).border = KENARLIK

    ornekler = [
        ["MERT TÜZE", "2025/203385", 2026, "Şubat", 7018.88, "ÖRNEK KAYIT"],
        ["MERT TÜZE", "2025/203385", 2026, "Mart", 7018.88, "ÖRNEK KAYIT"],
        ["MERT TÜZE", "2025/203385", 2026, "Nisan", 7018.88, "ÖRNEK KAYIT"],
        ["AHMET BAŞOĞLU", "2025/55555", 2026, "Mayıs", 6000.00, "ÖRNEK KAYIT"],
    ]
    for r, satir in enumerate(ornekler, start=2):
        for kol, deger in enumerate(satir, start=1):
            kes.cell(row=r, column=kol, value=deger)

    dv_personel = DataValidation(type="list", formula1="PersonelAdlari", allow_blank=True)
    dv_personel.errorStyle = "warning"
    kes.add_data_validation(dv_personel)
    dv_personel.add(f"A2:A{KESINTI_SON_SATIR}")

    dv_ay = DataValidation(type="list", formula1=f'"{AYLAR}"', allow_blank=True)
    kes.add_data_validation(dv_ay)
    dv_ay.add(f"D2:D{KESINTI_SON_SATIR}")
    return kes


def panel_sayfasi(wb):
    ws = wb.create_sheet("PANEL")
    ws.sheet_properties.tabColor = "FF0000"
    genislik(ws, [3, 30, 16, 16, 16])

    c = ws.cell(row=1, column=2, value="AY SONU KESİNTİ PANELİ")
    c.font = Font(bold=True, size=15, color=LACIVERT)

    ws.cell(row=3, column=2, value="Kesinti Yılı:").font = Font(bold=True)
    ws.cell(row=3, column=3, value=2026)
    ws.cell(row=4, column=2, value="Kesinti Ayı:").font = Font(bold=True)
    ws.cell(row=4, column=3, value="Haziran")
    for r in (3, 4):
        h = ws.cell(row=r, column=3)
        h.fill = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
        h.border = KENARLIK
        h.alignment = Alignment(horizontal="center")
        h.font = Font(bold=True, size=12)

    ws.cell(row=6, column=2, value="Bu ay kesilecek toplam (önizleme):").font = Font(bold=True)
    t = ws.cell(row=6, column=3)
    t.value = f"=SUM('İCRA DOSYALARI'!$N$2:$N${DOSYA_SON_SATIR})"
    t.number_format = PARA
    t.font = Font(bold=True, size=12, color="C00000")
    t.border = KENARLIK

    bilgi = ws.cell(row=8, column=2)
    bilgi.value = ("Yıl ve ayı seçin, sonra aşağıdaki BUTON ALANI'ndaki 'KESİNTİLERİ HESAPLA ve KAYDET' "
                   "butonuna basın. Makro kurulu değilse önce KULLANIM sayfasındaki 5 adımı yapın "
                   "(butonlar Alt+F8 > BUTONLARI_KUR ile gelir). Butonsuz kullanım: Alt+F8 > KESINTILERI_KAYDET.")
    bilgi.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B8:E9")

    ws.merge_cells("B11:E15")
    alan = ws.cell(row=11, column=2, value="BUTON ALANI\n(Alt+F8 > BUTONLARI_KUR çalıştırınca "
                                           "butonlar buraya eklenir)")
    alan.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alan.fill = PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")
    alan.font = Font(color="808080", italic=True)

    dv_yil = DataValidation(type="whole", operator="between", formula1="2020", formula2="2100")
    dv_yil.error = "2020-2100 arası bir yıl girin."
    ws.add_data_validation(dv_yil)
    dv_yil.add("C3")
    dv_ay = DataValidation(type="list", formula1=f'"{AYLAR}"')
    ws.add_data_validation(dv_ay)
    dv_ay.add("C4")
    return ws


def ozet_sayfasi(wb):
    ozet = wb.create_sheet("PERSONEL ÖZET")
    ozet.sheet_properties.tabColor = YESIL
    baslik_yaz(ozet, ["Personel Adı Soyadı", "Net Maaş (TL)", "Dosya Sayısı",
                      "Toplam İcra Tutarı (TL)", "Toplam Tahsil Edilen (TL)",
                      "Toplam Kalan Borç (TL)", "Ödemedeki Dosya No",
                      "Ödemedeki Dosyanın Kalanı (TL)", "Bu Ay Kesilecek (TL)"],
               otomatikler=(2, 3, 4, 5, 6, 7, 8, 9))
    genislik(ozet, [24, 13, 11, 17, 18, 17, 17, 18, 14])
    ozet.freeze_panes = "A2"

    d = f"'İCRA DOSYALARI'!$A$2:$A${DOSYA_SON_SATIR}"
    for r in range(2, OZET_SON_SATIR + 1):
        ozet.cell(row=r, column=2).value = (
            f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},PERSONEL!$A$2:$C${PERSONEL_SON},3,0),""))')
        ozet.cell(row=r, column=3).value = f'=IF($A{r}="","",COUNTIF({d},$A{r}))'
        for kol, kaynak in ((4, "$I"), (5, "$Q"), (6, "$R"), (9, "$N")):
            ozet.cell(row=r, column=kol).value = (
                f'=IF($A{r}="","",SUMIF({d},$A{r},'
                f"'İCRA DOSYALARI'!{kaynak}$2:{kaynak}${DOSYA_SON_SATIR}))")
        esle = f'MATCH($A{r}&"|ÖDEMEDE",\'İCRA DOSYALARI\'!$V$2:$V${DOSYA_SON_SATIR},0)'
        ozet.cell(row=r, column=7).value = (
            f'=IF($A{r}="","",IFERROR(INDEX(\'İCRA DOSYALARI\'!$E$2:$E${DOSYA_SON_SATIR},{esle}),"-"))')
        ozet.cell(row=r, column=8).value = (
            f'=IF($A{r}="","",IFERROR(INDEX(\'İCRA DOSYALARI\'!$R$2:$R${DOSYA_SON_SATIR},{esle}),"-"))')
        for kol in (2, 4, 5, 6, 8, 9):
            ozet.cell(row=r, column=kol).number_format = PARA
        for kol in range(1, 10):
            ozet.cell(row=r, column=kol).border = KENARLIK

    ozet.cell(row=2, column=1, value="MERT TÜZE")
    ozet.cell(row=3, column=1, value="AHMET BAŞOĞLU")
    return ozet


def main(yol):
    wb = Workbook()
    kullanim_sayfasi(wb)
    personel_sayfasi(wb)
    icra_sayfasi(wb)
    kesinti_sayfasi(wb)
    panel_sayfasi(wb)
    ozet_sayfasi(wb)

    wb.defined_names["PersonelAdlari"] = DefinedName(
        "PersonelAdlari", attr_text=f"PERSONEL!$A$2:$A${PERSONEL_SON}")

    wb.save(yol)
    print(f"Yazıldı: {yol}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "Personel_Icra_Takip.xlsx")
